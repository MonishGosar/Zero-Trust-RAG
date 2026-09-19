import json
import time

import pytest
from fastapi.testclient import TestClient

from backend.azure import extract_chart_from_markdown, sanitize_ui
from backend.config import Settings
from backend.main import create_app
from backend.models import Block
from backend.parsing import chunk_document, parse_document, parse_markdown
from tests.auth_helpers import sign_in
from tests.fakes import FakeAzure


@pytest.fixture
def setup(tmp_path):
    settings = Settings(_env_file=None, data_dir=tmp_path, retrieval_score_threshold=0.05)
    provider = FakeAzure()
    with TestClient(create_app(settings, provider)) as client:
        sign_in(client)
        yield client, provider


def upload_ready(client, name="report.md", text="# Revenue\n\nAPAC revenue was $82M."):
    response = client.post("/documents", files={"file": (name, text.encode())})
    assert response.status_code == 202, response.text
    document_id = response.json()["document_id"]
    for _ in range(100):
        detail = client.get(f"/documents/{document_id}").json()
        if detail["document"]["status"] in {"ready", "failed"}:
            return document_id, detail
        time.sleep(0.02)
    pytest.fail("Ingestion did not finish")


def test_upload_retrieve_cite_and_download(setup):
    client, provider = setup
    document_id, detail = upload_ready(client)
    assert detail["document"]["status"] == "ready"
    assert detail["chunks"][0]["section"] == "Revenue"
    retrieved = client.post("/retrieve", json={"question": "What was APAC revenue?"})
    assert retrieved.status_code == 200
    assert retrieved.json()[0]["document_id"] == document_id
    assert provider.answer_calls == 0
    response = client.post("/chat", json={"question": "What was APAC revenue?"})
    assert response.status_code == 200
    answer = response.json()
    assert answer["grounded"] is True
    assert answer["sources"][0]["document_id"] == document_id
    assert answer["sources"][0]["page"] is None
    assert provider.last_sources[0].text.endswith("APAC revenue was $82M.\n\n")
    assert client.get(f"/documents/{document_id}/file").content.startswith(b"# Revenue")


def test_table_atomic_and_heading_page_provenance():
    table = "| Region | Revenue |\n| --- | --- |\n| APAC | $82M |\n| Europe | $65M |"
    chunks = chunk_document(
        parse_markdown("# Finance\n\n" + table + "\n\nClosing note."),
        "doc",
        "report.md",
        target_tokens=10,
    )
    tables = [c for c in chunks if c.content_type == "table"]
    assert len(tables) == 1 and table in tables[0].text
    assert tables[0].section == "Finance"
    pages = chunk_document(
        [Block(text="Page one", page=1), Block(text="Page two", page=2)], "doc", "report.pdf"
    )
    assert [c.page for c in pages] == [1, 2]


def test_unicode_prose_preserved_and_bounded():
    text = "Knowledge café 東京 " * 1000
    chunks = chunk_document([Block(text=text)], "doc", "note.txt", target_tokens=80)
    assert all(c.token_count <= 82 for c in chunks)
    assert " ".join(c.text.strip() for c in chunks).split() == text.split()


def test_large_table_fails_explicitly():
    with pytest.raises(ValueError, match="table or section"):
        chunk_document([Block(text="APAC | 82 |\n" * 4000, content_type="table")], "d", "t.md")


def test_generative_ui_is_bounded_to_chart_and_table_vocabulary():
    chart = sanitize_ui({
        "root": {"component": "Chart", "props": {
            "variant": "bar", "data": [{"label": "Q1", "revenue": 42}], "citations": [1]
        }}
    })
    assert chart["root"]["component"] == "Chart"
    assert sanitize_ui({"root": {"component": "Script", "props": {}}}) is None
    assert sanitize_ui({"root": {"component": "Chart", "props": {
        "variant": "bar", "data": [{"label": "Q1", "revenue": "invented"}], "citations": [1]
    }}}) is None


def test_chart_fallback_extracts_cited_markdown_table():
    answer = """Source [2].
| Quarter | Months | Sales |
|---|---:|---:|
| Q1 | Apr-Jun | 350,832.00 |
| Q2 | Jul-Sep | 1,286,320.00 |
"""
    chart = extract_chart_from_markdown("Make a bar chart of sales", answer)
    assert chart["root"]["props"]["data"][0]["Sales"] == 350832.0
    assert chart["root"]["props"]["citations"] == [2]


def test_unknown_and_unready_documents_never_reach_model(setup):
    client, provider = setup
    response = client.post("/chat", json={"question": "Revenue?", "document_ids": ["missing"]})
    assert response.status_code == 409
    assert provider.answer_calls == 0
    provider.fail_embeddings = True
    _, detail = upload_ready(client)
    assert detail["document"]["status"] == "failed"
    assert "secret" not in detail["document"]["error"]
    provider.fail_embeddings = False
    assert client.post("/chat", json={"question": "Revenue?"}).json()["sources"] == []
    assert provider.answer_calls == 0


def test_scope_filters_before_generation(setup):
    client, provider = setup
    first, _ = upload_ready(client)
    second, _ = upload_ready(client, "second.md", "# Revenue\n\nAPAC revenue is confidential.")
    answer = client.post("/chat", json={"question": "APAC revenue", "document_ids": [first]}).json()
    assert answer["sources"][0]["document_id"] == first
    assert all(s.document_id != second for s in provider.last_sources)


@pytest.mark.parametrize(
    "answer",
    ["Revenue was $82M.", "Revenue was $82M. [99]", "Revenue was $82M. [1] Also $90M. [8]"],
)
def test_invalid_citations_suppress_answer(setup, answer):
    client, provider = setup
    upload_ready(client)
    provider.answer_text = answer
    result = client.post("/chat", json={"question": "APAC revenue?"}).json()
    assert result["grounded"] is False
    assert result["sources"] == []
    assert "$82M" not in result["answer"]


def test_empty_library_skips_azure_and_stream_finishes(setup):
    client, provider = setup
    result = client.post("/chat/stream", json={"question": "What is revenue?"})
    events = [json.loads(line) for line in result.text.splitlines()]
    assert [event["type"] for event in events] == ["status", "status", "answer"]
    assert events[-1]["sources"] == []
    assert provider.answer_calls == 0


def test_upload_validation_and_missing_config(setup, tmp_path):
    client, _ = setup
    assert client.post("/documents", files={"file": ("run.exe", b"abc")}).status_code == 415
    assert client.post("/documents", files={"file": ("empty.txt", b"")}).status_code == 400
    assert client.post("/documents", files={"file": ("fake.pdf", b"abc")}).status_code == 400
    assert client.post("/chat", json={"question": " "}).status_code == 422
    with TestClient(create_app(Settings(_env_file=None, data_dir=tmp_path / "unconfigured"))) as c:
        sign_in(c)
        assert not c.get("/health").json()["azure_configured"]
        assert c.post("/documents", files={"file": ("a.txt", b"hi")}).status_code == 503


def test_catalog_and_vectors_persist_across_restart(tmp_path):
    settings = Settings(_env_file=None, data_dir=tmp_path)
    with TestClient(create_app(settings, FakeAzure())) as client:
        sign_in(client)
        document_id, _ = upload_ready(client)
    with TestClient(create_app(settings, FakeAzure())) as client:
        sign_in(client)
        answer = client.post("/chat", json={"question": "APAC revenue"}).json()
        assert answer["sources"][0]["document_id"] == document_id


def test_docx_real_docling_conversion(tmp_path):
    from docx import Document

    doc = Document()
    doc.add_heading("Finance", level=1)
    doc.add_paragraph("APAC revenue was $82M.")
    table = doc.add_table(rows=2, cols=2)
    for cell, text in zip(
        [c for row in table.rows for c in row.cells], ["Region", "Revenue", "APAC", "$82M"]
    ):
        cell.text = text
    path = tmp_path / "finance.docx"
    doc.save(path)
    blocks, _ = parse_document(path)
    assert any(b.content_type == "table" and "$82M" in b.text for b in blocks)
    assert all(b.page is None for b in blocks)
    assert any(b.section == "Finance" for b in blocks)
